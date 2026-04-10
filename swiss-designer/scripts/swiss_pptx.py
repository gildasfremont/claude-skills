#!/usr/bin/env python3
"""
Swiss PPTX: helper library for International Typographic Style slide decks.

Usage:
    from swiss_pptx import SwissPPTX

    deck = SwissPPTX("output.pptx", title="My Deck", subtitle="March 2026",
                     author="Name", logo_path="logo.png")
    deck.title_slide(description="Short intro.")
    deck.section_slide(1, "Section Title", body="Body text.")
    deck.content_slide("More detail here.")
    deck.tip_slide("Tip label", "Tip body text.")
    deck.warn_slide("Warning label", "Warning text.")
    deck.list_slide("Steps", ["First step.", "Second step.", "Third step."])
    deck.closing_slide("Sign-off message.", "email@example.com")
    deck.save()
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

#Palette (grayscale)BLACK   = "#1A1A1A"
DARK    = "#2D2D2D"
BODY    = "#333333"
MUTED   = "#777777"
LIGHT   = "#BBBBBB"
ACCENT  = "#444444"
RULE    = "#D0D0D0"
TIP_BG  = "#F0F0F0"
WARN_BG = "#E8E8E8"
WHITE   = "#FFFFFF"

#FontFONT_NAME = "Liberation Sans"

def _rgb(hex_str):
    """Convert '#RRGGBB' to RGBColor."""
    return RGBColor.from_string(hex_str.lstrip("#"))

def _set_font(run, size, bold=False, italic=False, color=BODY):
    """Apply Swiss font styling to a text run."""
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)

def _add_textbox(slide, left, top, width, height):
    """Add a text box and return the shape."""
    return slide.shapes.add_textbox(left, top, width, height)

def _add_rect(slide, left, top, width, height, fill_hex, border=False):
    """Add a filled rectangle shape."""
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_hex)
    if not border:
        shape.line.fill.background()
    return shape


class SwissPPTX:
    """High-level API for building Swiss-style PPTX slide decks."""

    def __init__(self, outpath, title="Presentation", subtitle="", author="",
                 logo_path=None):
        self.outpath = outpath
        self.title_text = title
        self.subtitle_text = subtitle
        self.author = author
        self.logo_path = logo_path
        self.prs = Presentation()
        self.slide_num = 0

        # Dimensions
        self.SW = self.prs.slide_width
        self.SH = self.prs.slide_height

        # Grid
        self.MARGIN = Inches(0.6)
        self.LEFT = self.MARGIN
        self.RIGHT = self.SW - self.MARGIN
        self.TOP = self.MARGIN
        self.BOTTOM = self.SH - self.MARGIN
        self.COL_W = self.RIGHT - self.LEFT
        self.NUM_COL = Inches(1.2)
        self.TEXT_LEFT = self.LEFT + self.NUM_COL + Inches(0.15)
        self.TEXT_W = self.RIGHT - self.TEXT_LEFT

    def _blank_slide(self):
        """Add a blank slide and increment counter."""
        layout = self.prs.slide_layouts[6]  # Blank
        slide = self.prs.slides.add_slide(layout)
        self.slide_num += 1
        return slide

    def _draw_slide_number(self, slide):
        """Add slide number bottom-right."""
        tb = _add_textbox(slide,
                          self.RIGHT - Inches(0.8),
                          self.BOTTOM - Inches(0.1),
                          Inches(0.7), Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = str(self.slide_num)
        _set_font(run, 12, color=MUTED)

    def _draw_running_header(self, slide):
        """Add running header in top margin."""
        # Left: title
        tb = _add_textbox(slide, self.LEFT, Inches(0.15),
                          Inches(4), Inches(0.3))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = self.title_text
        _set_font(run, 10, color=LIGHT)

        # Right: subtitle
        if self.subtitle_text:
            tb2 = _add_textbox(slide, self.RIGHT - Inches(4), Inches(0.15),
                               Inches(4), Inches(0.3))
            p2 = tb2.text_frame.paragraphs[0]
            p2.alignment = PP_ALIGN.RIGHT
            run2 = p2.add_run()
            run2.text = self.subtitle_text
            _set_font(run2, 10, color=LIGHT)

    def _draw_accent_bar(self, slide):
        """Top accent bar, full width."""
        _add_rect(slide, 0, 0, self.SW, Inches(0.06), ACCENT)

    def _add_logo(self, slide, size=Inches(0.9)):
        """Place logo bottom-right."""
        lp = self.logo_path
        if lp and os.path.exists(lp):
            slide.shapes.add_picture(
                lp,
                self.RIGHT - size,
                self.BOTTOM - size + Inches(0.15),
                size, size
            )

    #Slide types
    def title_slide(self, description=""):
        """Create the title/cover slide."""
        slide = self._blank_slide()
        self._draw_accent_bar(slide)

        # Title
        y = Inches(1.0)
        tb = _add_textbox(slide, self.LEFT, y, self.COL_W, Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = self.title_text
        _set_font(run, 54, bold=True, color=BLACK)

        # Light rule
        y_rule = Inches(3.0)
        _add_rect(slide, self.LEFT, y_rule, self.COL_W, Inches(0.005), RULE)

        # Subtitle
        y_sub = y_rule + Inches(0.2)
        if self.subtitle_text:
            tb2 = _add_textbox(slide, self.LEFT, y_sub, self.COL_W, Inches(0.5))
            p2 = tb2.text_frame.paragraphs[0]
            run2 = p2.add_run()
            run2.text = self.subtitle_text
            _set_font(run2, 20, bold=True, color=MUTED)
            y_sub += Inches(0.4)

        # Author
        if self.author:
            tb3 = _add_textbox(slide, self.LEFT, y_sub, self.COL_W, Inches(0.4))
            p3 = tb3.text_frame.paragraphs[0]
            run3 = p3.add_run()
            run3.text = self.author
            _set_font(run3, 18, color=MUTED)
            y_sub += Inches(0.5)

        # Description
        if description:
            tb4 = _add_textbox(slide, self.LEFT, y_sub + Inches(0.3),
                               self.COL_W, Inches(1.2))
            tf4 = tb4.text_frame
            tf4.word_wrap = True
            p4 = tf4.paragraphs[0]
            run4 = p4.add_run()
            run4.text = description
            _set_font(run4, 16, color=MUTED)

        self._add_logo(slide)
        self._draw_slide_number(slide)

    def section_slide(self, num, title, body=""):
        """Slide with section number, title, and optional body text."""
        slide = self._blank_slide()
        self._draw_running_header(slide)

        y_top = Inches(1.4)

        # Section number in left column
        tb_num = _add_textbox(slide, self.LEFT, y_top - Inches(0.1),
                              self.NUM_COL, Inches(0.8))
        p_num = tb_num.text_frame.paragraphs[0]
        run_num = p_num.add_run()
        run_num.text = f"{num:02d}"
        _set_font(run_num, 48, bold=True, color=ACCENT)

        # Title in main column
        tb_title = _add_textbox(slide, self.TEXT_LEFT, y_top,
                                self.TEXT_W, Inches(1.2))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        run_title = p_title.add_run()
        run_title.text = title
        _set_font(run_title, 32, bold=True, color=BLACK)

        # Body text
        if body:
            tb_body = _add_textbox(slide, self.TEXT_LEFT, y_top + Inches(1.2),
                                   self.TEXT_W, Inches(3.5))
            tf_body = tb_body.text_frame
            tf_body.word_wrap = True
            p_body = tf_body.paragraphs[0]
            run_body = p_body.add_run()
            run_body.text = body
            _set_font(run_body, 18, color=BODY)

        self._draw_slide_number(slide)

    def content_slide(self, body, sub_heading=None):
        """Content slide within the current section (no section number)."""
        slide = self._blank_slide()
        self._draw_running_header(slide)

        y = Inches(1.4)

        if sub_heading:
            tb_sh = _add_textbox(slide, self.TEXT_LEFT, y,
                                 self.TEXT_W, Inches(0.6))
            p_sh = tb_sh.text_frame.paragraphs[0]
            run_sh = p_sh.add_run()
            run_sh.text = sub_heading
            _set_font(run_sh, 24, bold=True, color=DARK)
            y += Inches(0.7)

        tb = _add_textbox(slide, self.TEXT_LEFT, y, self.TEXT_W, Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = body
        _set_font(run, 18, color=BODY)

        self._draw_slide_number(slide)

    def tip_slide(self, label, body):
        """Slide with a tip callout box."""
        self._callout_slide(label, body, TIP_BG)

    def warn_slide(self, label, body):
        """Slide with a warning callout box."""
        self._callout_slide(label, body, WARN_BG)

    def _callout_slide(self, label, body, bg_hex):
        """Generic callout slide."""
        slide = self._blank_slide()
        self._draw_running_header(slide)

        y = Inches(1.4)
        box_pad = Inches(0.15)

        # Background rect
        _add_rect(slide,
                  self.TEXT_LEFT - box_pad,
                  y - box_pad,
                  self.TEXT_W + 2 * box_pad,
                  Inches(2.8),
                  bg_hex)

        # Label
        tb_label = _add_textbox(slide, self.TEXT_LEFT, y,
                                self.TEXT_W, Inches(0.5))
        p_label = tb_label.text_frame.paragraphs[0]
        run_label = p_label.add_run()
        run_label.text = label
        _set_font(run_label, 18, bold=True, color=DARK)

        # Body
        tb_body = _add_textbox(slide, self.TEXT_LEFT, y + Inches(0.55),
                               self.TEXT_W, Inches(2.0))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        p_body = tf_body.paragraphs[0]
        run_body = p_body.add_run()
        run_body.text = body
        _set_font(run_body, 16, color=BODY)

        self._draw_slide_number(slide)

    def list_slide(self, title, items):
        """Slide with a title and a numbered list."""
        slide = self._blank_slide()
        self._draw_running_header(slide)

        y = Inches(1.4)

        # Title
        tb_title = _add_textbox(slide, self.TEXT_LEFT, y,
                                self.TEXT_W, Inches(0.7))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        run_title = p_title.add_run()
        run_title.text = title
        _set_font(run_title, 28, bold=True, color=BLACK)

        # Items
        y_item = y + Inches(0.9)
        for i, item in enumerate(items):
            # Number
            tb_n = _add_textbox(slide, self.TEXT_LEFT, y_item,
                                Inches(0.4), Inches(0.4))
            p_n = tb_n.text_frame.paragraphs[0]
            run_n = p_n.add_run()
            run_n.text = str(i + 1)
            _set_font(run_n, 18, bold=True, color=ACCENT)

            # Text
            tb_t = _add_textbox(slide, self.TEXT_LEFT + Inches(0.45), y_item,
                                self.TEXT_W - Inches(0.45), Inches(0.5))
            tf_t = tb_t.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            run_t = p_t.add_run()
            run_t.text = item
            _set_font(run_t, 18, color=BODY)

            y_item += Inches(0.55)

        self._draw_slide_number(slide)

    def closing_slide(self, message, contact, logo_path=None):
        """Final slide with closing message and contact."""
        slide = self._blank_slide()
        self._draw_running_header(slide)

        y = Inches(2.5)

        # Light rule
        _add_rect(slide, self.LEFT, y, self.COL_W, Inches(0.004), RULE)
        y += Inches(0.35)

        # Message
        tb = _add_textbox(slide, self.LEFT, y, self.COL_W, Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = message
        _set_font(run, 18, italic=True, color=BODY)

        # Contact
        y += Inches(1.0)
        tb2 = _add_textbox(slide, self.LEFT, y, self.COL_W, Inches(0.4))
        p2 = tb2.text_frame.paragraphs[0]
        run2 = p2.add_run()
        run2.text = contact
        _set_font(run2, 14, color=MUTED)

        self._add_logo(slide, size=Inches(0.7))
        self._draw_slide_number(slide)

    def save(self):
        """Write the PPTX file."""
        self.prs.save(self.outpath)
        print(f"OK: {self.outpath}")
